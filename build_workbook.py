"""Builds Household_Finances.xlsx - a couples' budgeting & savings tracker.

Run:  python3 build_workbook.py
Re-running overwrites the workbook, so only do that to get a fresh blank copy.
"""
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Household_Finances.xlsx"

# ---------- sizes ----------
TX_FIRST, TX_LAST = 5, 2004          # Transactions rows
SV_FIRST, SV_LAST = 5, 1004          # Savings Log rows
RC_FIRST, RC_LAST = 5, 64            # Regular bills/income rows
ACC_FIRST, ACC_LAST = 8, 19          # Settings: account slots
INC_FIRST, INC_LAST = 8, 13          # Settings: income category slots
EXP_FIRST, EXP_LAST = 14, 38         # Settings: expense category slots
GOAL_FIRST, GOAL_LAST = 6, 17        # Savings goals

# ---------- styles ----------
FONT = "Arial"
GBP = '£#,##0.00;[Red]-£#,##0.00;"-"'
GBP0 = '£#,##0;[Red]-£#,##0;"-"'
PCT = '0.0%;[Red]-0.0%;"-"'
DATE_FMT = "dd/mm/yyyy"
MONTH_FMT = "mmm yyyy"

NAVY = "1F3864"
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")   # light yellow = type here
HELPER_FILL = PatternFill("solid", fgColor="EDEDED")  # grey = auto-calculated
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
SUB_FILL = PatternFill("solid", fgColor="D9E1F2")
TOTAL_FILL = PatternFill("solid", fgColor="E2EFDA")
thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

f_title = Font(name=FONT, size=16, bold=True, color=NAVY)
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_bold = Font(name=FONT, size=10, bold=True)
f_norm = Font(name=FONT, size=10)
f_input = Font(name=FONT, size=10, color="0000FF")
f_link = Font(name=FONT, size=10, color="008000")
f_note = Font(name=FONT, size=9, italic=True, color="595959")

wb = Workbook()


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = f_title
    ws["A2"] = sub
    ws["A2"].font = f_note
    ws.sheet_view.showGridLines = False


def header(ws, row, col, labels, helper_from=None):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=lab)
        c.font = f_head
        c.fill = HEAD_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
        if helper_from is not None and i >= helper_from:
            c.fill = PatternFill("solid", fgColor="595959")
    ws.row_dimensions[row].height = 30


def style(cell, font=f_norm, fmt=None, fill=None, bold=False, align=None):
    cell.font = Font(name=FONT, size=10, bold=True, color=font.color) if bold else font
    if fmt:
        cell.number_format = fmt
    if fill:
        cell.fill = fill
    if align:
        cell.alignment = Alignment(horizontal=align)
    cell.border = BORDER


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


def add_list_validation(ws, rng, formula, prompt=None):
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Pick a value from the drop-down list (edit the lists on the Settings tab)."
    dv.errorTitle = "Not in list"
    dv.errorStyle = "warning"
    if prompt:
        dv.prompt = prompt
        dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(rng)


# ======================================================================
# SETTINGS
# ======================================================================
st = wb.active
st.title = "Settings"
title(st, "Settings", "Set up names, accounts, categories and monthly budgets here. "
      "Yellow cells are yours to edit; everything else in the workbook updates automatically.")

st["B3"], st["B4"] = "Your name", "Partner's name"
st["C3"], st["C4"] = "Me", "Partner"
for r in (3, 4):
    style(st[f"B{r}"], f_bold)
    style(st[f"C{r}"], f_input, fill=INPUT_FILL)

# Owners list (feeds account owner drop-down)
header(st, 7, 5, ["Owners"])
st["E8"], st["E9"], st["E10"] = "=C3", "=C4", "Joint"
for r in (8, 9, 10):
    style(st[f"E{r}"], f_link if r < 10 else f_norm)

# Accounts
header(st, 7, 2, ["Account", "Owner"])
accounts = [
    ("Me - Current", "=C3"), ("Partner - Current", "=C4"), ("Joint Account", "Joint"),
    ("Me - Credit Card", "=C3"), ("Partner - Credit Card", "=C4"), ("Cash", "Joint"),
]
for i, r in enumerate(range(ACC_FIRST, ACC_LAST + 1)):
    a, o = accounts[i] if i < len(accounts) else (None, None)
    st.cell(row=r, column=2, value=a)
    st.cell(row=r, column=3, value=o)
    style(st.cell(row=r, column=2), f_input, fill=INPUT_FILL)
    style(st.cell(row=r, column=3), f_input, fill=INPUT_FILL)
add_list_validation(st, f"C{ACC_FIRST}:C{ACC_LAST}", "=$E$8:$E$10")
st[f"B{ACC_LAST + 2}"] = ("Owner decides whose column the account's money appears in. "
                          "Rename accounts freely - drop-downs update.")
st[f"B{ACC_LAST + 2}"].font = f_note

# Categories
header(st, 7, 8, ["Category", "Type", "Monthly budget (£)", "Essential?"])
income_cats = ["Salary", "Bonus", "Side Income", "Benefits", "Refunds", "Other Income"]
expense_cats = [  # (name, budget, essential) - example budgets, replace with your own
    ("Rent / Mortgage", 1200, "Yes"), ("Council Tax", 160, "Yes"), ("Energy", 140, "Yes"),
    ("Water", 40, "Yes"), ("Broadband & Phone", 70, "Yes"), ("Insurance", 60, "Yes"),
    ("Groceries", 450, "Yes"), ("Transport & Fuel", 180, "Yes"), ("Health", 30, "Yes"),
    ("Childcare & Kids", 0, "Yes"), ("Debt Repayments", 0, "Yes"),
    ("Eating Out", 150, "No"), ("Takeaways", 60, "No"), ("Subscriptions", 45, "No"),
    ("Entertainment", 80, "No"), ("Shopping & Clothes", 120, "No"),
    ("Personal Care", 40, "No"), ("Gifts", 40, "No"), ("Holidays", 100, "No"),
    ("Home & Garden", 50, "No"), ("Hobbies", 40, "No"), ("Other", 50, "No"),
]
for i, r in enumerate(range(INC_FIRST, INC_LAST + 1)):
    st.cell(row=r, column=8, value=income_cats[i] if i < len(income_cats) else None)
    st.cell(row=r, column=9, value="Income")
    st.cell(row=r, column=10, value=None)
    st.cell(row=r, column=11, value=None)
    style(st.cell(row=r, column=8), f_input, fill=INPUT_FILL)
    for c in (9, 10, 11):
        style(st.cell(row=r, column=c), fill=HELPER_FILL)
for i, r in enumerate(range(EXP_FIRST, EXP_LAST + 1)):
    n, b, e = expense_cats[i] if i < len(expense_cats) else (None, None, None)
    st.cell(row=r, column=8, value=n)
    st.cell(row=r, column=9, value="Expense")
    st.cell(row=r, column=10, value=b)
    st.cell(row=r, column=11, value=e)
    style(st.cell(row=r, column=8), f_input, fill=INPUT_FILL)
    style(st.cell(row=r, column=9), fill=HELPER_FILL)
    style(st.cell(row=r, column=10), f_input, GBP0, INPUT_FILL)
    style(st.cell(row=r, column=11), f_input, fill=INPUT_FILL, align="center")
add_list_validation(st, f"K{EXP_FIRST}:K{EXP_LAST}", '"Yes,No"')
st["J14"].comment = Comment("Example budgets - replace with your own monthly targets.", "Tracker")
st["K14"].comment = Comment("Essential = must-pay. 'No' marks discretionary spending - "
                            "the 'Where to Save' tab focuses on these.", "Tracker")
st[f"H{EXP_LAST + 2}"] = (f"Rows {INC_FIRST}-{INC_LAST} are income categories, rows "
                          f"{EXP_FIRST}-{EXP_LAST} expense categories. Fill blank rows to add more.")
st[f"H{EXP_LAST + 2}"].font = f_note

# Frequency factors (for Regular Bills)
header(st, 7, 13, ["Frequency", "Times / month"])
freqs = [("Weekly", "=52/12"), ("Fortnightly", "=26/12"), ("Monthly", 1),
         ("Quarterly", "=1/3"), ("Every 6 months", "=1/6"), ("Yearly", "=1/12")]
for i, (fq, fac) in enumerate(freqs):
    r = 8 + i
    st.cell(row=r, column=13, value=fq)
    st.cell(row=r, column=14, value=fac)
    style(st.cell(row=r, column=13))
    style(st.cell(row=r, column=14), fmt="0.000")

header(st, 7, 16, ["Transaction types"])
for i, t in enumerate(["Income", "Expense", "Transfer"]):
    st.cell(row=8 + i, column=16, value=t)
    style(st.cell(row=8 + i, column=16))
st["P12"] = "Transfer = moving money between your own accounts (e.g. into Joint)."
st["P12"].font = f_note

widths(st, {"A": 2, "B": 24, "C": 14, "D": 3, "E": 12, "F": 3, "G": 3, "H": 22, "I": 10,
            "J": 14, "K": 11, "L": 3, "M": 15, "N": 13, "O": 3, "P": 18})

# convenience refs
ACC_NAMES = f"Settings!$B${ACC_FIRST}:$B${ACC_LAST}"
ACC_OWNERS = f"Settings!$C${ACC_FIRST}:$C${ACC_LAST}"
CAT_ALL = f"Settings!$H${INC_FIRST}:$H${EXP_LAST}"
CAT_TYPE = f"Settings!$I${INC_FIRST}:$I${EXP_LAST}"
CAT_ESS = f"Settings!$K${INC_FIRST}:$K${EXP_LAST}"
FREQ_N, FREQ_F = "Settings!$M$8:$M$13", "Settings!$N$8:$N$13"

# ======================================================================
# TRANSACTIONS
# ======================================================================
tx = wb.create_sheet("Transactions")
title(tx, "Transactions - every incoming & outgoing",
      "One row per payment. Fill the yellow columns (A-H); grey columns fill themselves. "
      "Amounts are always positive - Type says which way the money went.")
labels = ["Date", "Description", "Account", "Type", "Category", "Amount (£)",
          "To account (transfers only)", "Notes", "Owner", "To owner", "Month", "Essential?"]
header(tx, 4, 1, labels, helper_from=8)
ex = date(2026, 10, 1)
samples = [  # example rows so the dashboard has something to show
    (date(2026, 9, 1), "Salary", "Me - Current", "Income", "Salary", 2650),
    (date(2026, 9, 1), "Salary", "Partner - Current", "Income", "Salary", 2400),
    (date(2026, 9, 1), "Rent", "Joint Account", "Expense", "Rent / Mortgage", 1200),
    (date(2026, 9, 10), "Sainsbury's", "Joint Account", "Expense", "Groceries", 470),
    (date(2026, 9, 18), "Deliveroo", "Partner - Current", "Expense", "Takeaways", 95),
    (date(2026, 9, 20), "Restaurant", "Me - Credit Card", "Expense", "Eating Out", 210),
    (date(2026, 10, 1), "Salary", "Me - Current", "Income", "Salary", 2650),
    (date(2026, 10, 1), "Salary", "Partner - Current", "Income", "Salary", 2400),
    (date(2026, 10, 1), "Monthly pot to joint", "Me - Current", "Transfer", None, 1100, "Joint Account"),
    (date(2026, 10, 1), "Monthly pot to joint", "Partner - Current", "Transfer", None, 1000, "Joint Account"),
    (date(2026, 10, 1), "Rent", "Joint Account", "Expense", "Rent / Mortgage", 1200),
    (date(2026, 10, 3), "Aldi big shop", "Joint Account", "Expense", "Groceries", 118.45),
    (date(2026, 10, 4), "Netflix", "Me - Credit Card", "Expense", "Subscriptions", 12.99),
    (date(2026, 10, 5), "Train to work", "Partner - Current", "Expense", "Transport & Fuel", 64.20),
    (date(2026, 10, 6), "Dinner out", "Me - Credit Card", "Expense", "Eating Out", 86.50),
    (date(2026, 10, 7), "New trainers", "Partner - Credit Card", "Expense", "Shopping & Clothes", 75),
]
for i, r in enumerate(range(TX_FIRST, TX_LAST + 1)):
    if i < len(samples):
        s = samples[i] + (None,) * (7 - len(samples[i]))
        for c, v in enumerate(s[:7]):
            tx.cell(row=r, column=c + 1 if c < 6 else 7, value=v)
        tx.cell(row=r, column=8, value="Example - delete")
    tx.cell(row=r, column=9, value=f'=IF(C{r}="","",IFERROR(INDEX({ACC_OWNERS},MATCH(C{r},{ACC_NAMES},0)),"?"))')
    tx.cell(row=r, column=10, value=f'=IF(G{r}="","",IFERROR(INDEX({ACC_OWNERS},MATCH(G{r},{ACC_NAMES},0)),"?"))')
    tx.cell(row=r, column=11, value=f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))')
    tx.cell(row=r, column=12, value=f'=IF(E{r}="","",IFERROR(INDEX({CAT_ESS},MATCH(E{r},{CAT_ALL},0))&"",""))')
    for c in range(1, 9):
        cell = tx.cell(row=r, column=c)
        cell.font = f_input
        cell.fill = INPUT_FILL
        cell.border = BORDER
    tx.cell(row=r, column=1).number_format = DATE_FMT
    tx.cell(row=r, column=6).number_format = GBP
    for c in range(9, 13):
        cell = tx.cell(row=r, column=c)
        cell.font = f_norm
        cell.fill = HELPER_FILL
        cell.border = BORDER
    tx.cell(row=r, column=11).number_format = MONTH_FMT
add_list_validation(tx, f"C{TX_FIRST}:C{TX_LAST}", f"={ACC_NAMES}", )
add_list_validation(tx, f"G{TX_FIRST}:G{TX_LAST}", f"={ACC_NAMES}")
add_list_validation(tx, f"D{TX_FIRST}:D{TX_LAST}", "=Settings!$P$8:$P$10")
add_list_validation(tx, f"E{TX_FIRST}:E{TX_LAST}", f"={CAT_ALL}")
dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
dv.error, dv.errorTitle = "Enter a positive amount - use Type to say income or expense.", "Positive amounts only"
tx.add_data_validation(dv)
dv.add(f"F{TX_FIRST}:F{TX_LAST}")
dv = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True)
dv.error, dv.errorTitle = "Enter a date, e.g. 14/10/2026.", "Date needed"
tx.add_data_validation(dv)
dv.add(f"A{TX_FIRST}:A{TX_LAST}")
# warn on missing category for expenses / missing destination for transfers
tx.conditional_formatting.add(
    f"E{TX_FIRST}:E{TX_LAST}",
    FormulaRule(formula=[f'AND(E{TX_FIRST}="",OR(D{TX_FIRST}="Expense",D{TX_FIRST}="Income"))'],
                fill=PatternFill("solid", fgColor="F8CBAD")))
tx.conditional_formatting.add(
    f"G{TX_FIRST}:G{TX_LAST}",
    FormulaRule(formula=[f'AND(G{TX_FIRST}="",D{TX_FIRST}="Transfer")'],
                fill=PatternFill("solid", fgColor="F8CBAD")))
tx.freeze_panes = "A5"
tx.auto_filter.ref = f"A4:L{TX_LAST}"
widths(tx, {"A": 12, "B": 26, "C": 22, "D": 11, "E": 20, "F": 13, "G": 22, "H": 20,
            "I": 11, "J": 11, "K": 11, "L": 11})

TX = lambda col: f"Transactions!${col}${TX_FIRST}:${col}${TX_LAST}"  # noqa: E731

# ======================================================================
# REGULAR BILLS & INCOME
# ======================================================================
rc = wb.create_sheet("Regular Bills")
title(rc, "Regular bills, subscriptions & income",
      "List anything that repeats (salary, rent, phone, gym, Netflix...). It is converted to a "
      "monthly & yearly cost so you can see what you're committed to - and what you could cancel.")
header(rc, 4, 1, ["Item", "Account", "Type", "Category", "Amount (£)", "Frequency",
                  "Per month", "Per year", "Owner", "Essential?", "Review notes"], helper_from=6)
rc_samples = [
    ("Salary", "Me - Current", "Income", "Salary", 2650, "Monthly", ""),
    ("Salary", "Partner - Current", "Income", "Salary", 2400, "Monthly", ""),
    ("Rent", "Joint Account", "Expense", "Rent / Mortgage", 1200, "Monthly", ""),
    ("Council tax", "Joint Account", "Expense", "Council Tax", 160, "Monthly", ""),
    ("Gas & electric", "Joint Account", "Expense", "Energy", 140, "Monthly", "Compare tariffs at renewal"),
    ("Broadband", "Joint Account", "Expense", "Broadband & Phone", 35, "Monthly", "Contract ends Mar - haggle"),
    ("Netflix", "Me - Credit Card", "Expense", "Subscriptions", 12.99, "Monthly", "Still watching?"),
    ("Gym", "Partner - Current", "Expense", "Hobbies", 9.25, "Weekly", ""),
    ("Car insurance", "Joint Account", "Expense", "Insurance", 540, "Yearly", ""),
    ("Amazon Prime", "Me - Credit Card", "Expense", "Subscriptions", 95, "Yearly", ""),
]
for i, r in enumerate(range(RC_FIRST, RC_LAST + 1)):
    if i < len(rc_samples):
        s = rc_samples[i]
        for c, v in enumerate(s[:6]):
            rc.cell(row=r, column=c + 1, value=v)
        rc.cell(row=r, column=11, value=s[6] or None)
    rc.cell(row=r, column=7, value=f'=IF(OR(E{r}="",F{r}=""),"",E{r}*IFERROR(INDEX({FREQ_F},MATCH(F{r},{FREQ_N},0)),0))')
    rc.cell(row=r, column=8, value=f'=IF(G{r}="","",G{r}*12)')
    rc.cell(row=r, column=9, value=f'=IF(B{r}="","",IFERROR(INDEX({ACC_OWNERS},MATCH(B{r},{ACC_NAMES},0)),"?"))')
    rc.cell(row=r, column=10, value=f'=IF(D{r}="","",IFERROR(INDEX({CAT_ESS},MATCH(D{r},{CAT_ALL},0))&"",""))')
    for c in (1, 2, 3, 4, 5, 6, 11):
        style(rc.cell(row=r, column=c), f_input, fill=INPUT_FILL)
    for c in (7, 8, 9, 10):
        style(rc.cell(row=r, column=c), fill=HELPER_FILL)
    for c in (5, 7, 8):
        rc.cell(row=r, column=c).number_format = GBP
add_list_validation(rc, f"B{RC_FIRST}:B{RC_LAST}", f"={ACC_NAMES}")
add_list_validation(rc, f"C{RC_FIRST}:C{RC_LAST}", '"Income,Expense"')
add_list_validation(rc, f"D{RC_FIRST}:D{RC_LAST}", f"={CAT_ALL}")
add_list_validation(rc, f"F{RC_FIRST}:F{RC_LAST}", f"={FREQ_N}")
rc.freeze_panes = "A5"
widths(rc, {"A": 22, "B": 22, "C": 10, "D": 20, "E": 12, "F": 15, "G": 12, "H": 12,
            "I": 11, "J": 11, "K": 30})
RC = lambda col: f"'Regular Bills'!${col}${RC_FIRST}:${col}${RC_LAST}"  # noqa: E731

# ======================================================================
# SAVINGS (goals) + SAVINGS LOG
# ======================================================================
sv = wb.create_sheet("Savings")
sl = wb.create_sheet("Savings Log")
title(sv, "Savings goals", "Set a goal, target and date. Log every deposit/withdrawal on the "
      "'Savings Log' tab - progress and the monthly amount needed update automatically.")
sv["A3"], sv["B3"] = "Today", "=TODAY()"
sv["A3"].font, sv["B3"].font, sv["B3"].number_format = f_bold, f_norm, DATE_FMT
header(sv, 5, 1, ["Goal", "Owner", "Target (£)", "Target date", "Starting balance (£)",
                  "Added since (£)", "Current balance (£)", "Still to save (£)", "% complete",
                  "Months left", "Needed per month (£)"], helper_from=5)
goals = [("Emergency fund", "Joint", 6000, date(2027, 6, 30), 1500),
         ("Holiday 2027", "Joint", 2500, date(2027, 7, 1), 200),
         ("New car", "=Settings!C3", 4000, date(2028, 1, 1), 0)]
SL = lambda col: f"'Savings Log'!${col}${SV_FIRST}:${col}${SV_LAST}"  # noqa: E731
for i, r in enumerate(range(GOAL_FIRST, GOAL_LAST + 1)):
    if i < len(goals):
        for c, v in enumerate(goals[i]):
            sv.cell(row=r, column=c + 1, value=v)
    sv.cell(row=r, column=6, value=f'=IF(A{r}="","",SUMIFS({SL("D")},{SL("B")},A{r}))')
    sv.cell(row=r, column=7, value=f'=IF(A{r}="","",E{r}+F{r})')
    sv.cell(row=r, column=8, value=f'=IF(A{r}="","",MAX(0,C{r}-G{r}))')
    sv.cell(row=r, column=9, value=f'=IF(OR(A{r}="",N(C{r})=0),"",MIN(1,G{r}/C{r}))')
    sv.cell(row=r, column=10, value=f'=IF(OR(A{r}="",D{r}=""),"",MAX(0,(YEAR(D{r})-YEAR($B$3))*12+MONTH(D{r})-MONTH($B$3)))')
    sv.cell(row=r, column=11, value=f'=IF(OR(A{r}="",D{r}=""),"",IF(H{r}=0,0,IF(J{r}=0,H{r},H{r}/J{r})))')
    for c in range(1, 6):
        style(sv.cell(row=r, column=c), f_input, fill=INPUT_FILL)
    for c in range(6, 12):
        style(sv.cell(row=r, column=c), fill=HELPER_FILL)
    for c in (3, 5, 6, 7, 8, 11):
        sv.cell(row=r, column=c).number_format = GBP
    sv.cell(row=r, column=4).number_format = DATE_FMT
    sv.cell(row=r, column=9).number_format = "0%"
add_list_validation(sv, f"B{GOAL_FIRST}:B{GOAL_LAST}", "=Settings!$E$8:$E$10")
tr_ = GOAL_LAST + 1
sv.cell(row=tr_, column=1, value="Total")
for c, col in ((3, "C"), (5, "E"), (6, "F"), (7, "G"), (8, "H"), (11, "K")):
    sv.cell(row=tr_, column=c, value=f"=SUM({col}{GOAL_FIRST}:{col}{GOAL_LAST})")
    style(sv.cell(row=tr_, column=c), fmt=GBP, fill=TOTAL_FILL, bold=True)
style(sv.cell(row=tr_, column=1), fill=TOTAL_FILL, bold=True)
sv.cell(row=tr_, column=9, value=f'=IF(C{tr_}=0,"",G{tr_}/C{tr_})')
style(sv.cell(row=tr_, column=9), fmt="0%", fill=TOTAL_FILL, bold=True)
sv.conditional_formatting.add(f"I{GOAL_FIRST}:I{GOAL_LAST}",
                              DataBarRule(start_type="num", start_value=0, end_type="num",
                                          end_value=1, color="70AD47"))
sv.freeze_panes = "A6"
widths(sv, {"A": 22, "B": 12, "C": 12, "D": 12, "E": 14, "F": 13, "G": 14, "H": 13,
            "I": 11, "J": 9, "K": 14})

title(sl, "Savings log", "Every time you put money into (or take it out of) savings, add a row. "
      "Use a negative amount for a withdrawal.")
header(sl, 4, 1, ["Date", "Goal", "From account", "Amount (£)", "Notes", "Owner", "Month"],
       helper_from=5)
sl_samples = [(date(2026, 8, 30), "Emergency fund", "Joint Account", 300),
              (date(2026, 9, 30), "Emergency fund", "Joint Account", 300),
              (date(2026, 10, 2), "Holiday 2027", "Me - Current", 150),
              (date(2026, 10, 2), "New car", "Me - Current", 200),
              (date(2026, 10, 3), "Holiday 2027", "Partner - Current", 150)]
for i, r in enumerate(range(SV_FIRST, SV_LAST + 1)):
    if i < len(sl_samples):
        for c, v in enumerate(sl_samples[i]):
            sl.cell(row=r, column=c + 1, value=v)
        sl.cell(row=r, column=5, value="Example - delete")
    sl.cell(row=r, column=6, value=f'=IF(C{r}="","",IFERROR(INDEX({ACC_OWNERS},MATCH(C{r},{ACC_NAMES},0)),"?"))')
    sl.cell(row=r, column=7, value=f'=IF(A{r}="","",DATE(YEAR(A{r}),MONTH(A{r}),1))')
    for c in range(1, 6):
        style(sl.cell(row=r, column=c), f_input, fill=INPUT_FILL)
    for c in (6, 7):
        style(sl.cell(row=r, column=c), fill=HELPER_FILL)
    sl.cell(row=r, column=1).number_format = DATE_FMT
    sl.cell(row=r, column=4).number_format = GBP
    sl.cell(row=r, column=7).number_format = MONTH_FMT
add_list_validation(sl, f"B{SV_FIRST}:B{SV_LAST}", f"=Savings!$A${GOAL_FIRST}:$A${GOAL_LAST}")
add_list_validation(sl, f"C{SV_FIRST}:C{SV_LAST}", f"={ACC_NAMES}")
sl.freeze_panes = "A5"
widths(sl, {"A": 12, "B": 22, "C": 22, "D": 13, "E": 24, "F": 11, "G": 11})

# ======================================================================
# DASHBOARD
# ======================================================================
db = wb.create_sheet("Dashboard", 0)
title(db, "Monthly dashboard", "Change the month in the yellow cell. Everything below "
      "recalculates from Transactions, Savings Log and Regular Bills.")
db["B3"] = "Month to show"
db["B3"].font = f_bold
db["C3"] = date(2026, 10, 1)
style(db["C3"], Font(name=FONT, size=11, bold=True, color="0000FF"), MONTH_FMT, INPUT_FILL)
db["C3"].comment = Comment("Type any date in the month you want, e.g. 01/11/2026.", "Tracker")
db["D3"] = '=DATE(YEAR(C3),MONTH(C3),1)'
db["D3"].number_format = MONTH_FMT
db["D3"].font = Font(name=FONT, size=10, color="FFFFFF")  # normalised month (hidden helper)
M = "$D$3"

# --- Section 1: who earned / spent / saved ---
db["B5"] = "Who earned, spent and saved"
db["B5"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
header(db, 6, 2, ["", "", "", "Joint", "Household"])
db["C6"], db["D6"] = "=Settings!C3", "=Settings!C4"
rows = ["Income", "Spending", "Transfers in", "Transfers out", "Saved",
        "Left over", "Savings rate", "Regular bills (planned / month)"]
for i, lab in enumerate(rows):
    r = 7 + i
    db.cell(row=r, column=2, value=lab)
    style(db.cell(row=r, column=2), bold=(lab in ("Left over",)))
    for j, col in enumerate("CDE"):
        own = f"{col}$6"
        f = {
            "Income": f'=SUMIFS({TX("F")},{TX("I")},{own},{TX("D")},"Income",{TX("K")},{M})',
            "Spending": f'=SUMIFS({TX("F")},{TX("I")},{own},{TX("D")},"Expense",{TX("K")},{M})',
            "Transfers in": f'=SUMIFS({TX("F")},{TX("J")},{own},{TX("D")},"Transfer",{TX("K")},{M})',
            "Transfers out": f'=SUMIFS({TX("F")},{TX("I")},{own},{TX("D")},"Transfer",{TX("K")},{M})',
            "Saved": f'=SUMIFS({SL("D")},{SL("F")},{own},{SL("G")},{M})',
            "Left over": f"={col}7-{col}8+{col}9-{col}10-{col}11",
            "Savings rate": f'=IF({col}7>0,{col}11/{col}7,"")',
            "Regular bills (planned / month)": f'=SUMIFS({RC("G")},{RC("I")},{own},{RC("C")},"Expense")',
        }[lab]
        db.cell(row=r, column=3 + j, value=f)
    db.cell(row=r, column=6, value=(f'=IF(F7>0,F11/F7,"")' if lab == "Savings rate"
                                    else f"=SUM(C{r}:E{r})"))
    for c in range(3, 7):
        cell = db.cell(row=r, column=c)
        style(cell, fmt=PCT if lab == "Savings rate" else GBP,
              fill=TOTAL_FILL if lab == "Left over" else None, bold=(lab == "Left over"))
db["B15"] = ("Transfers = money moved between your own accounts (e.g. into Joint); they cancel "
             "out in the Household column. Left over = income - spending + transfers in - out - saved.")
db["B15"].font = f_note

# --- Section 2: spending by category vs budget ---
db["B17"] = "Spending by category vs budget"
db["B17"].font = Font(name=FONT, size=12, bold=True, color=NAVY)
header(db, 18, 2, ["Category", "Budget", "", "", "Joint", "Total spent", "Left in budget",
                   "% of budget used", "3-month average"])
db["D18"], db["E18"] = "=Settings!C3", "=Settings!C4"
r0 = 19
n_exp = EXP_LAST - EXP_FIRST + 1
for i in range(n_exp):
    r, s = r0 + i, EXP_FIRST + i
    cat = f"$B{r}"
    db.cell(row=r, column=2, value=f'=IF(Settings!H{s}="","",Settings!H{s})')
    db.cell(row=r, column=3, value=f'=IF(B{r}="","",N(Settings!J{s}))')
    for j, col in enumerate("DEF"):
        db.cell(row=r, column=4 + j,
                value=f'=IF($B{r}="","",SUMIFS({TX("F")},{TX("E")},{cat},{TX("D")},"Expense",{TX("I")},{col}$18,{TX("K")},{M}))')
    db.cell(row=r, column=7, value=f'=IF(B{r}="","",SUM(D{r}:F{r}))')
    db.cell(row=r, column=8, value=f'=IF(B{r}="","",C{r}-G{r})')
    db.cell(row=r, column=9, value=f'=IF(OR(B{r}="",N(C{r})=0),"",G{r}/C{r})')
    db.cell(row=r, column=10, value=(
        f'=IF(B{r}="","",SUMIFS({TX("F")},{TX("E")},{cat},{TX("D")},"Expense",'
        f'{TX("K")},">="&EDATE({M},-2),{TX("K")},"<="&{M})/3)'))
    for c in range(2, 11):
        style(db.cell(row=r, column=c), fmt=PCT if c == 9 else (None if c == 2 else GBP))
rt = r0 + n_exp
db.cell(row=rt, column=2, value="Total")
for c in range(3, 11):
    L = get_column_letter(c)
    db.cell(row=rt, column=c, value=(f'=IF(C{rt}=0,"",G{rt}/C{rt})' if c == 9
                                     else f"=SUM({L}{r0}:{L}{rt - 1})"))
for c in range(2, 11):
    style(db.cell(row=rt, column=c), fmt=PCT if c == 9 else GBP, fill=TOTAL_FILL, bold=True)
rng = f"I{r0}:I{rt}"
db.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["1"],
                              fill=PatternFill("solid", fgColor="F8CBAD"),
                              font=Font(name=FONT, color="9C0006", bold=True)))
db.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.85", "1"],
                              fill=PatternFill("solid", fgColor="FFEB9C")))
db.conditional_formatting.add(f"H{r0}:H{rt}", CellIsRule(operator="lessThan", formula=["0"],
                              font=Font(name=FONT, color="C00000", bold=True)))
db.conditional_formatting.add(f"G{r0}:G{rt - 1}", DataBarRule(start_type="num", start_value=0,
                              end_type="max", color="5B9BD5"))
db.cell(row=rt + 1, column=2, value="Red = over budget, amber = 85%+ used. "
        "3-month average covers the selected month and the two before it.").font = f_note

# --- Section 3: year at a glance ---
yr = rt + 3
db.cell(row=yr, column=2, value="Year at a glance (household)").font = Font(name=FONT, size=12, bold=True, color=NAVY)
header(db, yr + 1, 2, ["Month", "Income", "Spending", "Saved", "Left over"])
for k in range(12):
    r = yr + 2 + k
    db.cell(row=r, column=2, value=f"=DATE(YEAR({M}),{k + 1},1)")
    db.cell(row=r, column=3, value=f'=SUMIFS({TX("F")},{TX("D")},"Income",{TX("K")},$B{r})')
    db.cell(row=r, column=4, value=f'=SUMIFS({TX("F")},{TX("D")},"Expense",{TX("K")},$B{r})')
    db.cell(row=r, column=5, value=f'=SUMIFS({SL("D")},{SL("G")},$B{r})')
    db.cell(row=r, column=6, value=f"=C{r}-D{r}-E{r}")
    for c in range(2, 7):
        style(db.cell(row=r, column=c), fmt=MONTH_FMT if c == 2 else GBP)
ye = yr + 14
db.cell(row=ye, column=2, value="Year total")
for c in range(3, 7):
    L = get_column_letter(c)
    db.cell(row=ye, column=c, value=f"=SUM({L}{yr + 2}:{L}{ye - 1})")
for c in range(2, 7):
    style(db.cell(row=ye, column=c), fmt=GBP, fill=TOTAL_FILL, bold=True)

# charts
ch = BarChart()
ch.type = "bar"
ch.title = "Budget vs spent (selected month)"
ch.style = 10
ch.add_data(Reference(db, min_col=3, min_row=18, max_row=r0 + len(expense_cats) - 1), titles_from_data=True)
ch.add_data(Reference(db, min_col=7, min_row=18, max_row=r0 + len(expense_cats) - 1), titles_from_data=True)
ch.set_categories(Reference(db, min_col=2, min_row=r0, max_row=r0 + len(expense_cats) - 1))
ch.y_axis.title = "£"
ch.y_axis.numFmt = "£#,##0"
ch.height, ch.width = 14, 16
db.add_chart(ch, "L17")

lc = BarChart()
lc.title = "Income vs spending by month"
lc.style = 10
lc.add_data(Reference(db, min_col=3, max_col=5, min_row=yr + 1, max_row=yr + 13), titles_from_data=True)
lc.set_categories(Reference(db, min_col=2, min_row=yr + 2, max_row=yr + 13))
lc.x_axis.number_format = "mmm"
lc.y_axis.numFmt = "£#,##0"
lc.height, lc.width = 9, 16
db.add_chart(lc, f"H{yr}")

db.freeze_panes = "A4"
widths(db, {"A": 2, "B": 30, "C": 13, "D": 13, "E": 13, "F": 13, "G": 13, "H": 14, "I": 12, "J": 14})

# ======================================================================
# WHERE TO SAVE
# ======================================================================
ws = wb.create_sheet("Where to Save", 1)
title(ws, "Where can we save?", "Uses the month picked on the Dashboard and your last 3 months "
      "of spending. Adjust the yellow cut % to see what trimming discretionary spending is worth.")
ws["B4"] = "Month analysed"
ws["C4"] = f"=Dashboard!{M}"
ws["C4"].number_format = MONTH_FMT
ws["B5"] = "Cut discretionary spending by"
ws["C5"] = 0.2
ws["C5"].comment = Comment("What-if lever: try 10%, 20%, 30%...", "Tracker")
style(ws["B4"], bold=True)
style(ws["C4"], f_link, MONTH_FMT)
style(ws["B5"], bold=True)
style(ws["C5"], f_input, "0%", INPUT_FILL)

header(ws, 7, 2, ["Quick stats", "Value"])
TOTAL_ROW = rt  # dashboard total row
stats = [
    ("Discretionary ('non-essential') spend this month",
     f'=SUMIFS({TX("F")},{TX("L")},"No",{TX("D")},"Expense",{TX("K")},Dashboard!$D$3)', GBP),
    ("...as % of all spending this month", f'=IF(Dashboard!G{TOTAL_ROW}=0,"",C8/Dashboard!G{TOTAL_ROW})', PCT),
    ("Categories over budget this month", f'=COUNTIF(Dashboard!I{r0}:I{TOTAL_ROW - 1},">1")', "0"),
    ("Subscriptions - yearly cost", f'=SUMIFS({RC("H")},{RC("D")},"Subscriptions",{RC("C")},"Expense")', GBP),
    ("All non-essential regular bills - yearly cost", f'=SUMIFS({RC("H")},{RC("J")},"No",{RC("C")},"Expense")', GBP),
    ("Spending with no category (fix these!)",
     f'=COUNTIFS({TX("D")},"Expense",{TX("E")},"")', "0"),
    ("Potential yearly saving at the cut % above", f"=SUM(H16:H{15 + n_exp})", GBP),
]
for i, (lab, f, fmt) in enumerate(stats):
    r = 8 + i
    ws.cell(row=r, column=2, value=lab)
    ws.cell(row=r, column=3, value=f)
    style(ws.cell(row=r, column=2))
    style(ws.cell(row=r, column=3), fmt=fmt, bold=(i == len(stats) - 1),
          fill=TOTAL_FILL if i == len(stats) - 1 else None)

header(ws, 15, 2, ["Category", "Budget / month", "3-month avg", "Essential?",
                   "Avg vs budget", "Avg x 12 (yearly)", "Saving / year at cut %", "Rank helper"])
for i in range(n_exp):
    r, dr = 16 + i, r0 + i
    ws.cell(row=r, column=2, value=f"=Dashboard!B{dr}")
    ws.cell(row=r, column=3, value=f"=Dashboard!C{dr}")
    ws.cell(row=r, column=4, value=f"=Dashboard!J{dr}")
    ws.cell(row=r, column=5, value=f'=IF(B{r}="","",Settings!K{EXP_FIRST + i})')
    ws.cell(row=r, column=6, value=f'=IF(OR(B{r}="",N(C{r})=0),"",D{r}-C{r})')
    ws.cell(row=r, column=7, value=f'=IF(B{r}="","",D{r}*12)')
    ws.cell(row=r, column=8, value=f'=IF(E{r}="No",G{r}*$C$5,"")')
    ws.cell(row=r, column=9, value=f'=IF(AND(E{r}="No",N(D{r})>0),D{r}+ROW()/1000000,-1)')
    for c in range(2, 10):
        style(ws.cell(row=r, column=c), f_link if c in (2, 3, 4, 5) else f_norm,
              fmt=None if c in (2, 5, 9) else GBP)
    ws.cell(row=r, column=5).alignment = Alignment(horizontal="center")
    ws.cell(row=r, column=9).font = Font(name=FONT, size=8, color="A6A6A6")
ws.column_dimensions["I"].hidden = True
ws.conditional_formatting.add(f"F16:F{15 + n_exp}", CellIsRule(
    operator="greaterThan", formula=["0"], fill=PatternFill("solid", fgColor="F8CBAD")))
ws.cell(row=16 + n_exp, column=2,
        value="Pink 'Avg vs budget' = you're regularly spending more than planned there.").font = f_note

header(ws, 7, 6, ["Biggest discretionary categories (3-month avg)", "", "Avg / month"])
ws.merge_cells("F7:G7")
for k in range(5):
    r = 8 + k
    hr = f"$I$16:$I${15 + n_exp}"
    ws.cell(row=r, column=6, value=f'=IF(LARGE({hr},{k + 1})<=0,"",INDEX($B$16:$B${15 + n_exp},MATCH(LARGE({hr},{k + 1}),{hr},0)))')
    ws.merge_cells(start_row=r, start_column=6, end_row=r, end_column=7)
    ws.cell(row=r, column=8, value=f'=IF(F{r}="","",INDEX($D$16:$D${15 + n_exp},MATCH(F{r},$B$16:$B${15 + n_exp},0)))')
    style(ws.cell(row=r, column=6), bold=True)
    style(ws.cell(row=r, column=7))
    style(ws.cell(row=r, column=8), fmt=GBP)
ws.cell(row=13, column=6, value="Start here: these are the easiest places to cut back.").font = f_note

widths(ws, {"A": 2, "B": 46, "C": 15, "D": 14, "E": 11, "F": 14, "G": 16, "H": 18})

# ======================================================================
# HOW TO USE
# ======================================================================
hw = wb.create_sheet("How to use", 0)
title(hw, "Household Finances - how to use", "Couples' budget, spending and savings tracker")
steps = [
    ("Colour key", ""),
    ("  Yellow cells", "Type here - these are your inputs."),
    ("  Grey cells", "Calculated automatically - don't type over them."),
    ("  Blue text / green text", "Blue = something you entered, green = pulled from another tab."),
    ("", ""),
    ("1. Settings", "Put in both names, your accounts (each one owned by you, your partner or Joint), "
                    "your categories and a monthly budget for each. Mark each expense Essential Yes/No."),
    ("2. Regular Bills", "List everything that repeats - salaries, rent, bills, subscriptions. Shows "
                         "your committed monthly/yearly cost and flags things worth reviewing."),
    ("3. Transactions", "Log each incoming/outgoing as it happens (or paste from a bank export). "
                        "Pick account, type and category from the drop-downs; amounts are positive."),
    ("", "Moving money between your own accounts (e.g. your pot into Joint)? Use Type = Transfer "
         "and fill 'To account' - it won't be counted as spending."),
    ("4. Savings / Savings Log", "Set goals with targets and dates; log every deposit. See % done "
                                 "and how much you need to put away each month."),
    ("5. Dashboard", "Pick a month. See income, spending, savings and what's left over for each of "
                     "you, Joint and the household - plus every category against its budget."),
    ("6. Where to Save", "Your biggest discretionary categories, where you're over budget, the yearly "
                         "cost of subscriptions and what a % cut would save over a year."),
    ("", ""),
    ("Example data", "Rows marked 'Example - delete' show the format. Clear the yellow cells on "
                     "Transactions, Regular Bills, Savings and Savings Log when you start for real "
                     "(don't delete whole rows - just clear the contents)."),
    ("Room", f"Transactions: {TX_LAST - TX_FIRST + 1} rows, Savings Log: {SV_LAST - SV_FIRST + 1}, "
             f"Regular Bills: {RC_LAST - RC_FIRST + 1}. Start a new copy each year."),
]
for i, (a, b) in enumerate(steps):
    r = 4 + i
    hw.cell(row=r, column=2, value=a).font = f_bold
    hw.cell(row=r, column=3, value=b).font = f_norm
    hw.cell(row=r, column=3).alignment = Alignment(wrap_text=True, vertical="top")
    hw.cell(row=r, column=2).alignment = Alignment(vertical="top")
hw["B5"].fill = INPUT_FILL
hw["B6"].fill = HELPER_FILL
hw["B7"].font = Font(name=FONT, size=10, bold=True, color="0000FF")
widths(hw, {"A": 2, "B": 26, "C": 100})

# default fonts for any unstyled cells
for sheet in wb.worksheets:
    sheet.sheet_properties.tabColor = {
        "How to use": "808080", "Dashboard": NAVY, "Where to Save": "C00000",
        "Transactions": "FFC000", "Regular Bills": "FFC000", "Savings": "70AD47",
        "Savings Log": "70AD47", "Settings": "808080"}[sheet.title]
    for row in sheet.iter_rows():
        for c in row:
            if c.value is not None and c.font.name != FONT:
                c.font = Font(name=FONT, size=10, bold=c.font.b, italic=c.font.i,
                              color=c.font.color)

wb.active = wb.sheetnames.index("Dashboard")
wb.save(OUT)
print("wrote", OUT)
